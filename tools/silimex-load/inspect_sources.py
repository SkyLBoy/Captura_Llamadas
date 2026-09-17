import openpyxl, json, pathlib, collections, datetime, hashlib
root=pathlib.Path(__file__).parent
sources={
 'abechuco12': pathlib.Path('C:/Users/VNCAdmin-12/Documents/SILIMEX_2026/SILIMEXMONSE/SILIMEX_ABECHUCO_2026_120926.xlsm'),
 'agosto': pathlib.Path('C:/Users/VNCAdmin-12/Documents/SILIMEX_2026/SILIMEX_AGOSTO_2026/SILIMEX_AGOSTO_2026.xlsx'),
 'ordaz': pathlib.Path('C:/Users/VNCAdmin-12/Documents/SILIMEX_2026/SILIMEXCARLOS/SILIMEX_-_ORDAZ_-_2026.xlsm'),
 'abechuco': pathlib.Path('C:/Users/VNCAdmin-12/Documents/SILIMEX_2026/SILIMEXMONSE/SILIMEX_ABECHUCO_2026_150926.xlsm')}
for agent,path in sources.items():
 book=openpyxl.load_workbook(path,read_only=True,data_only=True)
 report={'file':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheetNames':book.sheetnames,'sheets':[]}
 extracted={}
 for sheet in book:
  if sheet.title.strip().upper()=='ENCUESTAS': continue
  if agent!='agosto' and sheet.title.strip().upper() not in ['LLAMADAS','BASE','ENCUESTAS 2','DATOS','HOJA 3','HOJA3']: continue
  rows=[]
  for n,row in enumerate(sheet.iter_rows(values_only=True),1):
   if any(v is not None for v in row): rows.append({'row':n,'values':list(row)})
  extracted[sheet.title]=rows
  report['sheets'].append({'name':sheet.title,'rows':len(rows),'columns':sheet.max_column,'sample':rows[:3]})
 (root/f'{agent}-data.json').write_text(json.dumps(extracted,ensure_ascii=False,default=str),encoding='utf-8')
 (root/f'{agent}-inspection.json').write_text(json.dumps(report,ensure_ascii=False,default=str,indent=2),encoding='utf-8')
 print(json.dumps(report,ensure_ascii=False,default=str))
